"""Daily Sales Entry <-> .xlsx.

Layout (one sheet, "Daily Sales Entry"):

    A: section / label   B: field   C: Value (input)   D: Computed   F: field key

Import reads **only column C keyed by column F** - cosmetic columns A/B/D may be
edited freely. Fill the Value column; do not delete the key column or reorder.
Per SDD ADR-5 the parsed payload is always recomputed; sheet totals are only
compared, never trusted.
"""

from __future__ import annotations

import io
import math
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from svr_backend.calc.amounts import is_blank
from svr_backend.calc.daily_sales_entry import OIL_KEYS, OIL_LABELS, compute_payload

SHEET = "Daily Sales Entry"
_MAX_OILS = 8  # 5 fixed + up to 3 operator-added
_MAX_ROWS = 10  # expenses / cards / new-credits / old-credits template rows
_EPS = 0.01

_Row = tuple[str, str, Any, Any, str | None]


def _r(label: str, field: str, value: Any = None, computed: Any = None,
       key: str | None = None) -> _Row:
    return (label, field, value, computed, key)


def _section(name: str) -> _Row:
    return (name, "", None, None, None)


# --------------------------------------------------------------------------- build


def _gas_rows(payload: dict, result: dict) -> list[_Row]:
    rows: list[_Row] = [_section("1. Gas Sale(s)")]
    for fuel, name in (("hs", "Diesel (HS)"), ("ms", "Petrol (MS)")):
        g = payload.get(fuel) or {}
        rg = result.get(fuel) or {}
        rows += [
            _r(name, "Current Reading", g.get("current"), key=f"{fuel}.current"),
            _r(name, "Last Shift Reading", g.get("last"), key=f"{fuel}.last"),
            _r(name, "Rate Per Pump", g.get("rate"), key=f"{fuel}.rate"),
            _r(name, "Consumption", computed=rg.get("cons"), key=f"_chk.{fuel}.cons"),
            _r(name, "Amount", computed=rg.get("amount"), key=f"_chk.{fuel}.amount"),
        ]
    rows.append(_r("", "Gas Total Amt", computed=result.get("gas_total"), key="_chk.gas_total"))
    return rows


def _oil_rows(payload: dict, result: dict) -> list[_Row]:
    oils = payload.get("oils") or []
    r_oils = result.get("oils") or []
    rows: list[_Row] = [_section("2. Oil Sale(s)")]
    for i in range(_MAX_OILS):
        o = oils[i] if i < len(oils) else {}
        ro = r_oils[i] if i < len(r_oils) else {}
        key = OIL_KEYS[i] if i < len(OIL_KEYS) else ""
        label = o.get("label") or OIL_LABELS.get(key, f"Oil {i + 1}")
        rows += [
            _r(label, "Quantity", o.get("qty"), key=f"oils.{i}.qty"),
            _r(label, "Rate", o.get("rate"), key=f"oils.{i}.rate"),
            _r(label, "Before Stock", o.get("opening"), key=f"oils.{i}.opening"),
            _r(label, "After Sale Stock", computed=ro.get("closing"), key=f"_chk.oils.{i}.closing"),
            _r(label, "Amount", computed=ro.get("amount"), key=f"_chk.oils.{i}.amount"),
        ]
    rows.append(_r("", "Total Amt Oil(s)", computed=result.get("oil_total"), key="_chk.oil_total"))
    return rows


def _list_section(title: str, label: str, values: list, base_key: str, total_key: str,
                  total_label: str, total_val: Any) -> list[_Row]:
    rows: list[_Row] = [_section(title)]
    for i in range(max(_MAX_ROWS, len(values))):
        v = values[i] if i < len(values) else None
        rows.append(_r(label, f"Row {i + 1}", v, key=f"{base_key}.{i}"))
    rows.append(_r("", total_label, computed=total_val, key=f"_chk.{total_key}"))
    return rows


def _new_credit_rows(payload: dict, result: dict) -> list[_Row]:
    nc = payload.get("new_credits") or []
    r_nc = result.get("new_credit_amounts") or []
    rows: list[_Row] = [_section("5. Today New Credit(s)")]
    for i in range(max(_MAX_ROWS, len(nc))):
        row = nc[i] if i < len(nc) else {}
        amt = r_nc[i] if i < len(r_nc) else None
        n = f"Row {i + 1}"
        rows += [
            _r("New credit", f"{n} - In Ltrs", row.get("ltrs"), key=f"new_credits.{i}.ltrs"),
            _r("New credit", f"{n} - Rate", row.get("rate"), key=f"new_credits.{i}.rate"),
            _r("New credit", f"{n} - Amount", computed=amt, key=f"_chk.new_credit_amounts.{i}"),
        ]
    rows.append(
        _r("", "Total Amt New Credits", computed=result.get("new_credits_total"),
           key="_chk.new_credits_total")
    )
    return rows


def _summary_rows(payload: dict, result: dict) -> list[_Row]:
    s = "Summary"
    return [
        _section("7. Summary - Cash Hand Off"),
        _r(s, "Cash (Gas+Oils) Total Amt", computed=result.get("sum_cash"), key="_chk.sum_cash"),
        _r(s, "Expenses Total Amt", computed=result.get("sum_expenses"), key="_chk.sum_expenses"),
        _r(s, "Phone Pay Settled Total Amt", payload.get("phone_pay_settled"),
           key="phone_pay_settled"),
        _r(s, "Phone Pay Not Settled Total Amt", payload.get("phone_pay_unsettled"),
           key="phone_pay_unsettled"),
        _r(s, "New Credits Total Amt", computed=result.get("sum_new_credits"),
           key="_chk.sum_new_credits"),
        _r(s, "Credit Cards Swiping Total Amt", computed=result.get("sum_credit_cards"),
           key="_chk.sum_credit_cards"),
        _r(s, "Night Cash Hand Off Total Amt", payload.get("night_cash"), key="night_cash"),
        _r(s, "Net Bal Hand Off", computed=result.get("net_bal_hand_off"),
           key="_chk.net_bal_hand_off"),
    ]


def _all_rows(payload: dict, result: dict, meta: dict) -> list[_Row]:
    rows: list[_Row] = [
        _r("Header", "Date", meta.get("shift_date"), key="meta.shift_date"),
        _r("Header", "Pump", meta.get("pump_serial"), key="meta.pump_serial"),
        _r("Header", "Submitted By", meta.get("submitted_by"), key="_info.submitted_by"),
        _r("Header", "Entry Mode", meta.get("entry_mode"), key="_info.entry_mode"),
    ]
    rows += _gas_rows(payload, result)
    rows += _oil_rows(payload, result)
    rows += _list_section(
        "3. Expenses", "Expense", payload.get("expenses") or [], "expenses",
        "expenses_total", "Total Amt Expenses", result.get("expenses_total"),
    )
    rows += _list_section(
        "4. Credit Cards Swiping(s)", "Card swipe", payload.get("credit_card_amounts") or [],
        "credit_card_amounts", "credit_cards_total", "Total Amt Credit Cards",
        result.get("credit_cards_total"),
    )
    rows += _new_credit_rows(payload, result)
    rows += _list_section(
        "6. Old/Pending Credit Received", "Old credit", payload.get("old_credit_amounts") or [],
        "old_credit_amounts", "old_credit_total", "Old Credit Total (reference only)",
        result.get("old_credit_total"),
    )
    rows += _summary_rows(payload, result)
    return rows


_SECTION_FILL = PatternFill("solid", fgColor="1F3B5C")
_HEADER_FILL = PatternFill("solid", fgColor="EEF1F5")


def build_workbook(record: dict) -> bytes:
    """`record` = {shift_date, pump_serial, submitted_by, entry_mode, payload, result}."""
    payload = record.get("payload") or {}
    result = record.get("result") or compute_payload(payload)
    meta = {k: record.get(k) for k in ("shift_date", "pump_serial", "submitted_by", "entry_mode")}

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws["A1"] = "SVR Indian Oil Service Station - Daily Sales Report (data export)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (
        "Fill the Value column. Do not delete column F (field key) or reorder rows. "
        "Totals recompute on import."
    )
    ws["A2"].font = Font(italic=True, size=9, color="666666")

    hdr = 4
    titles = ("Section / Label", "Field", "Value", "Computed", "", "field key (do not edit)")
    for col, name in enumerate(titles, 1):
        c = ws.cell(row=hdr, column=col, value=name)
        c.font = Font(bold=True)
        c.fill = _HEADER_FILL

    r = hdr + 1
    for label, field_name, value, computed, key in _all_rows(payload, result, meta):
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=field_name)
        if value is not None:
            ws.cell(row=r, column=3, value=value)
        if computed is not None:
            ws.cell(row=r, column=4, value=computed)
        if key:
            ws.cell(row=r, column=6, value=key)
        if field_name == "" and key is None:  # section header row
            for col in range(1, 7):
                cc = ws.cell(row=r, column=col)
                cc.font = Font(bold=True, color="FFFFFF")
                cc.fill = _SECTION_FILL
        r += 1

    for col, w in {"A": 34, "B": 30, "C": 16, "D": 14, "E": 2, "F": 26}.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=hdr + 1, min_col=3, max_col=4):
        for cell in row:
            cell.alignment = Alignment(horizontal="right")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def blank_template(pump_serial: str | None = None) -> bytes:
    empty = {
        "hs": {}, "ms": {},
        "oils": [{"label": OIL_LABELS[k]} for k in OIL_KEYS],
        "expenses": [], "credit_card_amounts": [], "new_credits": [], "old_credit_amounts": [],
    }
    return build_workbook({
        "shift_date": None, "pump_serial": pump_serial, "submitted_by": None,
        "entry_mode": "excel", "payload": empty, "result": compute_payload(empty),
    })


# --------------------------------------------------------------------------- parse


_KEY_PREFIXES = (
    "hs.", "ms.", "oils.", "expenses.", "credit_card_amounts.",
    "new_credits.", "old_credit_amounts.", "meta.", "_chk.", "_info.",
)
_KEY_EXACT = ("phone_pay_settled", "phone_pay_unsettled", "night_cash")


def _looks_like_field_key(v: Any) -> bool:
    """True only for our own dotted/exact key vocabulary (column F) - NOT just any
    non-empty string. openpyxl pads every row out to the sheet's widest row, so a
    natural workbook's own header text (e.g. "Amount", "Signature") can otherwise
    land in column F by accident and be mistaken for one of our field keys."""
    return isinstance(v, str) and (v in _KEY_EXACT or v.startswith(_KEY_PREFIXES))


def _num_or_str(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        # A lone "-" is the paper form's own "nothing to report" marker (every
        # real client sample uses it consistently) - treated as blank, same as
        # calc.amounts.is_blank (2026-09-11).
        return None if s in ("", "-") else s
    return v


def _blank_payload() -> dict:
    return {
        "hs": {}, "ms": {},
        "oils": [{} for _ in range(_MAX_OILS)],
        "expenses": [None] * _MAX_ROWS,
        "credit_card_amounts": [None] * _MAX_ROWS,
        "new_credits": [{} for _ in range(_MAX_ROWS)],
        "old_credit_amounts": [None] * _MAX_ROWS,
    }


def _assign(payload: dict, key: str, value: Any) -> None:
    parts = key.split(".")
    head = parts[0]
    if head in ("hs", "ms") and len(parts) == 2:
        payload[head][parts[1]] = value
    elif head == "oils" and len(parts) == 3:
        payload["oils"][int(parts[1])][parts[2]] = value
    elif head == "new_credits" and len(parts) == 3:
        payload["new_credits"][int(parts[1])][parts[2]] = value
    elif head in ("expenses", "credit_card_amounts", "old_credit_amounts") and len(parts) == 2:
        payload[head][int(parts[1])] = value
    elif head in ("phone_pay_settled", "phone_pay_unsettled", "night_cash"):
        payload[head] = value


# --------------------------------------------------------- paper-layout (fallback)
#
# Reads a workbook shaped like the *physical* Daily Sales Report - labels in early
# columns, values beside them - rather than our own keyed export/template. This is
# exactly what a person gets when they ask an AI chat tool (or anyone else) to type
# up a photographed handwritten form: a natural spreadsheet, not our internal
# format. Matched by exact cell text (not OCR), so it's reliable whenever the
# layout resembles the printed form; ADR-5 review still applies before Save.

_PUNCT_RE = re.compile(r"[^\w\s]")
_WORD_DIGIT_BOUNDARY_RE = re.compile(r"(?<=[a-z])(?=\d)|(?<=\d)(?=[a-z])")


def _txt(v: Any) -> str:
    """Lowercased, with punctuation folded to spaces, a word/number boundary
    inserted where letters and digits touch directly, and whitespace collapsed
    - tolerant of real-world label variance ("2T/1.20 ML" vs "2T-1.20ML" vs
    "Total1Lts" vs extra spaces) when matching a label/anchor. Applied
    consistently to both sides of every comparison below (including the hint
    constants, via ``_norm_hints``), so it only ever widens a match, never
    narrows one. Never used for extracting actual cell values - only for
    finding rows/columns.
    """
    if v is None:
        return ""
    s = _PUNCT_RE.sub(" ", str(v).lower())
    s = _WORD_DIGIT_BOUNDARY_RE.sub(" ", s)
    return " ".join(s.split())


def _row_txt(row: tuple) -> str:
    return " ".join(_txt(v) for v in row if v is not None)


def _norm_hints(hints: tuple[str, ...]) -> tuple[str, ...]:
    return tuple(_txt(h) for h in hints)


_HS_LABEL_HINTS = _norm_hints(("diesel", "hs-nz", "(hs"))
_MS_LABEL_HINTS = _norm_hints(("petrol", "ms-nz", "(ms"))

# Distinctive substrings per fixed oil row - first token alone is ambiguous between
# the two Acid Water rows, so each hint pins down the row uniquely.
_OIL_MATCH_HINTS: dict[str, tuple[str, ...]] = {
    "oil1": _norm_hints(("2t/1.20", "2t 1.20", "1.20 ml")),
    "oil2": _norm_hints(("2t/2.40", "2t 2.40", "2.40 ml")),
    "oil3": _norm_hints(("acid water total 1", "acid water 1 lt")),
    "oil4": _norm_hints(("acid water total 5", "acid water 5 lt")),
    "oil5": _norm_hints(("20/40", "20 40 engine")),
}


def _find_row(rows: list[tuple], *needles: str, start: int = 0) -> int | None:
    """First row index >= start whose cells together contain every needle."""
    needles_n = [_txt(n) for n in needles]
    for i in range(start, len(rows)):
        t = _row_txt(rows[i])
        if all(n in t for n in needles_n):
            return i
    return None


def _col_of(row: tuple, *needles: str) -> int | None:
    """First column index in ``row`` whose own cell text contains every needle."""
    needles_n = [_txt(n) for n in needles]
    for j, v in enumerate(row):
        cell = _txt(v)
        if all(n in cell for n in needles_n):
            return j
    return None


def _cell(row: tuple | None, col: int | None) -> Any:
    if row is None or col is None or col >= len(row):
        return None
    return _num_or_str(row[col])


def _rightmost_value(row: tuple, skip_first: int = 1) -> Any:
    """Rightmost non-blank cell, skipping the leading label column(s) - for simple
    Description | ... | Amount rows with no separate header to anchor a column on."""
    for v in reversed(row[skip_first:]):
        out = _num_or_str(v)
        if out is not None:
            return out
    return None


def _section_span(rows: list[tuple], title: str, next_title: str | None) -> tuple[int, int] | None:
    start = _find_row(rows, title)
    if start is None:
        return None
    if next_title is None:
        return start, len(rows)
    end = _find_row(rows, next_title, start=start + 1)
    return start, end if end is not None else len(rows)


def _parse_paper_layout(rows: list[tuple]) -> tuple[dict, dict, list[str]]:
    warnings = [
        "Read as a paper-form layout (no SVR field keys found) - this is a "
        "best-effort match on the form's own labels. Check EVERY value against "
        "the original form before Save (SDD ADR-5)."
    ]
    payload = _blank_payload()
    meta: dict[str, Any] = {"shift_date": None, "pump_serial": None}

    for row in rows[:4]:  # header block: pick up the first date-like cell
        for v in row:
            if hasattr(v, "strftime"):
                meta["shift_date"] = _fmt_date(v)
                break
        if meta["shift_date"]:
            break

    # ---- 1. Gas Sale(s) ----
    gas_hdr = _find_row(rows, "current reading", "last shift")
    if gas_hdr is None:
        warnings.append(
            "Could not find the Gas Sale(s) table - Diesel/Petrol readings were not read."
        )
    else:
        col_current = _col_of(rows[gas_hdr], "current reading")
        col_last = _col_of(rows[gas_hdr], "last shift")
        stop = _find_row(rows, "total amt", start=gas_hdr + 1)
        stop = stop if stop is not None else min(gas_hdr + 6, len(rows))
        for i in range(gas_hdr + 1, stop):
            label = _txt(rows[i][0] if rows[i] else None)
            if not label:
                continue
            if any(h in label for h in _HS_LABEL_HINTS):
                payload["hs"]["current"] = _cell(rows[i], col_current)
                payload["hs"]["last"] = _cell(rows[i], col_last)
            elif any(h in label for h in _MS_LABEL_HINTS):
                payload["ms"]["current"] = _cell(rows[i], col_current)
                payload["ms"]["last"] = _cell(rows[i], col_last)

    # ---- 2. Oil Sale(s) ----
    oil_hdr = _find_row(rows, "quantity", start=gas_hdr + 1 if gas_hdr is not None else 0)
    if oil_hdr is None:
        warnings.append("Could not find the Oil Sale(s) table - quantities were not read.")
    else:
        col_qty = _col_of(rows[oil_hdr], "quantity")
        stop = _find_row(rows, "total amt", start=oil_hdr + 1)
        stop = stop if stop is not None else min(oil_hdr + 8, len(rows))
        oils = []
        for key in OIL_KEYS:
            hints = _OIL_MATCH_HINTS[key]
            ridx = next(
                (i for i in range(oil_hdr + 1, stop) if any(h in _row_txt(rows[i]) for h in hints)),
                None,
            )
            oils.append({"qty": _cell(rows[ridx] if ridx is not None else None, col_qty)})
        payload["oils"] = oils

    # ---- 3. Expenses (3 fixed description rows, Amount is the last cell) ----
    exp_span = _section_span(rows, "expenses", "credit cards swiping")
    if exp_span:
        s, e = exp_span
        expenses: list[Any] = [None, None, None]
        for needle, idx in (
            ("daily diesel", 0), ("any other expenses", 1), ("night cash hand-off", 2),
        ):
            ridx = _find_row(rows, needle, start=s + 1)
            if ridx is not None and ridx < e:
                expenses[idx] = _rightmost_value(rows[ridx])
        payload["expenses"] = expenses

    # ---- 4. Credit Cards Swiping(s) ----
    cc_span = _section_span(rows, "credit cards swiping", "today new credit")
    if cc_span:
        s, e = cc_span
        hdr = _find_row(rows, "terminal id", start=s)
        if hdr is not None:
            payload["credit_card_amounts"] = [
                v
                for i in range(hdr + 1, e)
                if _row_txt(rows[i]) and "total amt" not in _row_txt(rows[i])
                for v in [_rightmost_value(rows[i], skip_first=0)]
                if v is not None
            ]

    # ---- 5. Today New Credit(s) ----
    nc_span = _section_span(rows, "today new credit", "old/pending credit")
    if nc_span:
        s, e = nc_span
        hdr = _find_row(rows, "in ltrs", start=s)
        if hdr is not None:
            col_ltrs = _col_of(rows[hdr], "in ltrs")
            col_rate = _col_of(rows[hdr], "rate")
            new_credits = []
            for i in range(hdr + 1, e):
                t = _row_txt(rows[i])
                if not t or "total amt" in t:
                    continue
                ltrs, rate = _cell(rows[i], col_ltrs), _cell(rows[i], col_rate)
                if ltrs is not None or rate is not None:
                    new_credits.append({"ltrs": ltrs, "rate": rate})
            payload["new_credits"] = new_credits

    # ---- 6. Old/Pending Credit Received (reference only) ----
    oc_span = _section_span(rows, "old/pending credit", "summary - cash hand off")
    if oc_span:
        s, e = oc_span
        hdr = _find_row(rows, "customer name", start=s)
        if hdr is not None:
            col_amt = _col_of(rows[hdr], "amount")
            payload["old_credit_amounts"] = [
                v for i in range(hdr + 1, e) if (v := _cell(rows[i], col_amt)) is not None
            ]

    # ---- 7. Summary - only the 3 operator-entered lines. Everything else here is
    #         computed and gets recomputed downstream; the sheet's own totals (Cash,
    #         Net Bal, ...) are never read, per ADR-5.
    summary_start = _find_row(rows, "summary - cash hand off")
    if summary_start is not None:
        for needle, key in (
            ("phone pay settled", "phone_pay_settled"),
            ("phone pay not settled", "phone_pay_unsettled"),
            ("night cash", "night_cash"),
        ):
            ridx = _find_row(rows, needle, start=summary_start + 1)
            if ridx is not None:
                payload[key] = _rightmost_value(rows[ridx])

    payload["oils"] = _trim_oils(payload["oils"])
    for k in ("expenses", "credit_card_amounts", "old_credit_amounts"):
        payload[k] = _rtrim(list(payload[k]))
    payload["new_credits"] = _rtrim_dicts(payload["new_credits"], ("ltrs", "rate"))

    return payload, meta, warnings


def _select_sheet(wb, pump_serial: str | None):
    """Pick the right sheet - always ours by name if present, else the one
    matching the selected Pump Serial Number in a multi-pump workbook (e.g. one
    file with a "Road ..." sheet and an "Office ..." sheet), else the active
    sheet if there's only one. Returns ``(worksheet, warning_or_None)``.

    Everything here is keyed by Pump Serial Number, never by sheet position -
    a workbook that happens to have the wrong sheet active must not silently
    import the wrong pump's data.
    """
    if SHEET in wb.sheetnames:
        return wb[SHEET], None
    names = wb.sheetnames
    if len(names) == 1:
        return wb.active, None
    if pump_serial:
        needle = pump_serial.strip().upper()
        for name in names:
            if needle and needle in name.upper():
                return wb[name], None
    return (
        wb.active,
        f"This workbook has {len(names)} sheets ({', '.join(names)}) - could not "
        f"match the selected Pump Serial Number to one of them, so '{wb.active.title}' "
        "was read. Select the correct Pump Serial Number and re-import if that's wrong.",
    )


def parse_workbook(data: bytes, pump_serial: str | None = None) -> tuple[dict, dict, list[str]]:
    """bytes -> (payload, meta, warnings). Recompute downstream; never trust sheet totals.

    ``pump_serial`` is the Pump Serial Number currently selected on the form -
    used only to pick the right sheet out of a multi-pump workbook; the parsed
    payload never trusts an in-file pump serial for identity either way.
    """
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws, sheet_warning = _select_sheet(wb, pump_serial)
    all_rows = list(ws.iter_rows(values_only=True))  # materialize once (read-only sheet)

    kv: dict[str, Any] = {}
    checks: dict[str, Any] = {}
    for row in all_rows:
        key = row[5] if len(row) > 5 else None
        if not _looks_like_field_key(key):
            continue
        if key.startswith("_chk."):
            checks[key[5:]] = _num_or_str(row[3])
        elif key.startswith("_info."):
            continue
        else:
            kv[key] = _num_or_str(row[2])

    if not kv and not checks:
        # Not one of our own exports/templates - most likely a natural, paper-shaped
        # workbook (e.g. a handwritten form typed up externally). Read it by matching
        # the physical form's own labels instead (exact cell text, not OCR, so it's
        # reliable whenever the layout resembles the printed form). ADR-5 review
        # still applies downstream - nothing here is trusted without a human Save.
        payload, meta, warnings = _parse_paper_layout(all_rows)
        if sheet_warning:
            warnings.insert(0, sheet_warning)
        return payload, meta, warnings

    warnings: list[str] = [sheet_warning] if sheet_warning else []
    payload = _blank_payload()
    for key, value in kv.items():
        if not key.startswith("meta."):
            _assign(payload, key, value)

    payload["oils"] = _trim_oils(payload["oils"])
    for k in ("expenses", "credit_card_amounts", "old_credit_amounts"):
        payload[k] = _rtrim(list(payload[k]))
    payload["new_credits"] = _rtrim_dicts(payload["new_credits"], ("ltrs", "rate"))

    meta = {
        "shift_date": _fmt_date(kv.get("meta.shift_date")),
        "pump_serial": str(kv["meta.pump_serial"]).strip() if kv.get("meta.pump_serial") else None,
    }

    _flag_mismatches(compute_payload(payload), checks, warnings)
    return payload, meta, warnings


def _trim_oils(oils: list[dict]) -> list[dict]:
    kept = list(oils[: len(OIL_KEYS)])  # always keep the 5 fixed rows
    for extra in oils[len(OIL_KEYS):]:
        if any(not is_blank(extra.get(f)) for f in ("qty", "rate", "opening")):
            kept.append(extra)
    return kept


def _rtrim(xs: list) -> list:
    while xs and is_blank(xs[-1]):
        xs.pop()
    return xs


def _rtrim_dicts(rows: list[dict], fields: tuple[str, ...]) -> list[dict]:
    while rows and all(is_blank(rows[-1].get(f)) for f in fields):
        rows.pop()
    return rows


def _fmt_date(v: Any) -> str | None:
    if v is None or v == "":
        return None
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v).strip() or None


def _look(tree: Any, dotted: str) -> Any:
    cur = tree
    for p in dotted.split("."):
        if isinstance(cur, list):
            cur = cur[int(p)] if p.isdigit() and int(p) < len(cur) else None
        elif isinstance(cur, dict):
            cur = cur.get(p)
        else:
            return None
    return cur


def _flag_mismatches(recomputed: dict, checks: dict, warnings: list[str]) -> None:
    for key, sheet_val in checks.items():
        if is_blank(sheet_val):
            continue
        want = _look(recomputed, key)
        try:
            if want is not None and not math.isclose(float(sheet_val), float(want), abs_tol=_EPS):
                warnings.append(
                    f"{key}: sheet shows {sheet_val}, recomputes to {want} - using recomputed."
                )
        except (TypeError, ValueError):
            continue
