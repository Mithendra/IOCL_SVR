"""Daily Sales Entry Excel import/export (BRD Section 5.5.1).

Standard data export - fields/values, not a visual replica of the printed form
(the layout-replication question is BRD-flagged, out of scope). Import never
saves and never trusts sheet totals (SDD ADR-5).
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from svr_backend.calc.daily_sales_entry import compute_payload
from svr_backend.excel import blank_template, build_workbook, parse_workbook

PUMP = "12BC4523V-RD"

_PAYLOAD = {
    "hs": {"current": 1317.52, "last": 1300, "rate": 105.36},
    "ms": {"current": 1000, "last": 900, "rate": 117.7},
    "oils": [{"label": "2T/1.20 ML", "qty": 4, "rate": 60, "opening": 100}],
    "expenses": ["500+100=600", 75],
    "credit_card_amounts": [1000, 250],
    "new_credits": [{"ltrs": 10, "rate": 105.36}],
    "old_credit_amounts": [300],
    "phone_pay_settled": 10,
    "phone_pay_unsettled": 20,
    "night_cash": 5000,
}


def _record(payload):
    return {
        "shift_date": "2026-08-12", "pump_serial": PUMP, "submitted_by": "gsales",
        "entry_mode": "manual", "payload": payload, "result": compute_payload(payload),
    }


# ------------------------------------------------------------------- unit: round-trip


def test_export_then_parse_preserves_inputs():
    data = build_workbook(_record(_PAYLOAD))
    payload, meta, warnings = parse_workbook(data)

    assert warnings == []
    assert meta == {"shift_date": "2026-08-12", "pump_serial": PUMP}
    assert payload["hs"] == {"current": 1317.52, "last": 1300, "rate": 105.36}
    assert payload["ms"]["rate"] == 117.7
    assert payload["oils"][0] == {"qty": 4, "rate": 60, "opening": 100}
    assert len(payload["oils"]) == 5  # 5 fixed rows kept, no phantom extras
    assert payload["expenses"] == ["500+100=600", 75]   # free-text expression survives
    assert payload["credit_card_amounts"] == [1000, 250]
    assert payload["new_credits"] == [{"ltrs": 10, "rate": 105.36}]
    assert payload["old_credit_amounts"] == [300]
    assert payload["phone_pay_unsettled"] == 20
    assert payload["night_cash"] == 5000


def test_recomputed_totals_match_after_round_trip():
    payload, _, _ = parse_workbook(build_workbook(_record(_PAYLOAD)))
    assert compute_payload(payload)["net_bal_hand_off"] == compute_payload(_PAYLOAD)["net_bal_hand_off"]


def test_blank_template_parses_clean():
    payload, meta, warnings = parse_workbook(blank_template(PUMP))
    assert warnings == []
    assert meta["pump_serial"] == PUMP
    assert payload["expenses"] == [] and payload["new_credits"] == []


def test_tampered_total_is_flagged_not_trusted():
    wb = load_workbook(io.BytesIO(build_workbook(_record(_PAYLOAD))))
    ws = wb["Daily Sales Entry"]
    for row in ws.iter_rows(min_col=1, max_col=6):
        if row[5].value == "_chk.net_bal_hand_off":
            row[3].value = 999999          # someone edited the computed cell
    buf = io.BytesIO()
    wb.save(buf)

    payload, _, warnings = parse_workbook(buf.getvalue())
    assert any("net_bal_hand_off" in w for w in warnings)
    # the payload still recomputes to the real figure - the sheet total is ignored
    assert compute_payload(payload)["net_bal_hand_off"] != 999999


# ------------------------------------------------------------------------------- API


def _create_entry(client, auth_headers):
    body = {
        "pump_serial": PUMP, "shift_date": "2026-08-12",
        "hs": {"current": "1317.52"}, "ms": {"current": "1000"},
        "oils": [{"qty": "4"}], "expenses": ["500+100=600"],
        "credit_card_amounts": ["1000"], "new_credits": [{"ltrs": "10", "rate": "105.36"}],
        "night_cash": "5000",
    }
    r = client.post("/daily-sales-entry", json=body, headers=auth_headers("Sales"))
    assert r.status_code == 201, r.text
    return r.json()["id"]


def test_export_endpoint_returns_xlsx(client, auth_headers):
    entry_id = _create_entry(client, auth_headers)
    r = client.get(f"/daily-sales-entry/{entry_id}/export-excel", headers=auth_headers("Sales"))
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in r.headers["content-disposition"]
    assert r.content[:2] == b"PK"  # xlsx is a zip


def test_export_missing_entry_404(client, auth_headers):
    assert client.get("/daily-sales-entry/9999/export-excel", headers=auth_headers("Sales")).status_code == 404


def test_template_endpoint(client, auth_headers):
    r = client.get("/daily-sales-entry/import-excel/template?pump_serial=X-OFF", headers=auth_headers("Sales"))
    assert r.status_code == 200
    assert r.content[:2] == b"PK"
    payload, meta, warnings = parse_workbook(r.content)
    assert meta["pump_serial"] == "X-OFF" and warnings == []


def test_import_endpoint_parses_and_recomputes(client, auth_headers):
    xlsx = build_workbook(_record(_PAYLOAD))
    r = client.post(
        "/daily-sales-entry/import-excel",
        files={"file": ("day.xlsx", xlsx, "application/octet-stream")},
        headers=auth_headers("Manager"),
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["warnings"] == []
    assert body["meta"]["pump_serial"] == PUMP
    assert body["payload"]["hs"]["current"] == 1317.52
    assert body["result"]["net_bal_hand_off"] == compute_payload(_PAYLOAD)["net_bal_hand_off"]


def test_import_rejects_non_spreadsheet(client, auth_headers):
    r = client.post(
        "/daily-sales-entry/import-excel",
        files={"file": ("x.xlsx", b"not a zip", "application/octet-stream")},
        headers=auth_headers("Sales"),
    )
    assert r.status_code == 400


def test_excel_endpoints_require_auth(client):
    assert client.get("/daily-sales-entry/1/export-excel").status_code == 401
    assert client.get("/daily-sales-entry/import-excel/template").status_code == 401
    assert client.post("/daily-sales-entry/import-excel").status_code == 401
