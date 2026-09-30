"""Fuel Procurement (client, 2026-09-30): RBAC, add/search/edit a load row,
posting to Monthly Expenses under the pre-existing "Fuel Load - HS"/"Fuel Load
- MS" categories with amount = IOCL Load x Rate (not Received litres), a
posted row going read-only, and the Excel export.
"""

from __future__ import annotations

import io

import openpyxl


def _add(client, headers, **overrides):
    body = {
        "shift_date": "2026-09-29",
        "fuel_type": "HS",
        "iocl_load_litres": 12000,
        "received_litres": 11970,
        "rate": 92.50,
        "vehicle_ref": "AP-16-TB-4519",
    }
    body.update(overrides)
    return client.post("/fuel-procurement", json=body, headers=headers)


def test_sales_has_no_access(client, auth_headers):
    assert _add(client, auth_headers("Sales")).status_code == 403
    assert client.get("/fuel-procurement", headers=auth_headers("Sales")).status_code == 403


def test_manager_can_add_a_load_and_lost_amount_are_computed(client, auth_headers):
    r = _add(client, auth_headers("Manager"))
    assert r.status_code == 201
    body = r.json()
    assert body["status"] == "draft"
    assert body["lost_litres"] == 30          # 12000 - 11970
    assert body["amount"] == 1110000.0        # 12000 x 92.50, off IOCL Load not Received


def test_a_typical_load_is_two_rows_same_date_and_ref(client, auth_headers):
    h = auth_headers("Owner")
    _add(client, h, fuel_type="HS", iocl_load_litres=12000, received_litres=11970, rate=92.50)
    _add(client, h, fuel_type="MS", iocl_load_litres=8000, received_litres=7982, rate=101.60)
    rows = client.get("/fuel-procurement?start=2026-09-29&end=2026-09-29", headers=h).json()
    assert len(rows) == 2
    assert {r["fuel_type"] for r in rows} == {"HS", "MS"}
    assert all(r["vehicle_ref"] == "AP-16-TB-4519" for r in rows)


def test_search_filters_by_date_fuel_and_status(client, auth_headers):
    h = auth_headers("Manager")
    _add(client, h, shift_date="2026-09-14", fuel_type="HS")
    _add(client, h, shift_date="2026-09-29", fuel_type="MS")

    sep = client.get("/fuel-procurement?start=2026-09-01&end=2026-09-20", headers=h).json()
    assert len(sep) == 1
    assert sep[0]["shift_date"] == "2026-09-14"

    hs_only = client.get("/fuel-procurement?fuel_type=HS", headers=h).json()
    assert len(hs_only) == 1
    assert hs_only[0]["fuel_type"] == "HS"

    not_posted = client.get("/fuel-procurement?status=draft", headers=h).json()
    assert len(not_posted) == 2


def test_save_edits_an_unposted_row(client, auth_headers):
    h = auth_headers("Manager")
    load_id = _add(client, h).json()["id"]
    r = client.put(
        f"/fuel-procurement/{load_id}",
        json={"received_litres": 11950, "rate": 93.00},
        headers=h,
    )
    assert r.status_code == 200
    body = r.json()
    assert body["received_litres"] == 11950
    assert body["lost_litres"] == 50          # 12000 - 11950
    assert body["amount"] == 1116000.0        # 12000 x 93.00


def test_post_creates_the_right_monthly_expense_line_and_locks_the_row(
    client, auth_headers, conn
):
    h = auth_headers("Owner")
    load_id = _add(client, h).json()["id"]

    r = client.post(f"/fuel-procurement/{load_id}/post", headers=h)
    assert r.status_code == 200
    posted = r.json()
    assert posted["status"] == "posted"
    assert posted["posted_expense_id"]

    expense = conn.execute(
        "SELECT me.amount, ec.name FROM monthly_expense me "
        "JOIN expense_category ec ON ec.id = me.category_id WHERE me.id = ?",
        (posted["posted_expense_id"],),
    ).fetchone()
    assert expense["name"] == "Fuel Load - HS"
    assert expense["amount"] == 1110000.0

    # A posted row is read-only.
    edit = client.put(
        f"/fuel-procurement/{load_id}", json={"rate": 999}, headers=h
    )
    assert edit.status_code == 409
    post_again = client.post(f"/fuel-procurement/{load_id}/post", headers=h)
    assert post_again.status_code == 409

    audit = conn.execute(
        "SELECT COUNT(*) c FROM audit_log WHERE table_name = 'fuel_load' "
        "AND action = 'update'"
    ).fetchone()["c"]
    assert audit >= 1


def test_ms_fuel_posts_to_the_ms_category(client, auth_headers, conn):
    h = auth_headers("Manager")
    load_id = _add(client, h, fuel_type="MS", rate=101.60).json()["id"]
    posted = client.post(f"/fuel-procurement/{load_id}/post", headers=h).json()
    expense = conn.execute(
        "SELECT ec.name FROM monthly_expense me "
        "JOIN expense_category ec ON ec.id = me.category_id WHERE me.id = ?",
        (posted["posted_expense_id"],),
    ).fetchone()
    assert expense["name"] == "Fuel Load - MS"


def test_export_excel_returns_a_readable_workbook(client, auth_headers):
    h = auth_headers("Manager")
    _add(client, h, fuel_type="HS", iocl_load_litres=12000, rate=92.50)
    _add(client, h, fuel_type="MS", iocl_load_litres=8000, rate=101.60)

    r = client.get("/fuel-procurement/export-excel?start=2026-09-01&end=2026-09-30", headers=h)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert "IOCL Load (L)" in values
    assert "HS" in values and "MS" in values


def test_export_monthly_excel_aggregates_by_calendar_month(client, auth_headers):
    h = auth_headers("Owner")
    _add(client, h, shift_date="2026-09-14", fuel_type="HS", iocl_load_litres=10000, rate=91.85)
    _add(client, h, shift_date="2026-09-14", fuel_type="MS", iocl_load_litres=10000, rate=101.05)
    _add(client, h, shift_date="2026-10-02", fuel_type="HS", iocl_load_litres=12000, rate=92.00)

    r = client.get("/fuel-procurement/export-monthly-excel", headers=h)
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    values = [cell.value for row in ws.iter_rows() for cell in row]
    assert "Sep 2026" in values
    assert "Oct 2026" in values
    # September's HS and MS rows both land on ONE month line, 10,000 L each.
    assert values.count(10000.0) == 2
