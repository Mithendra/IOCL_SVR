"""Extract every formula from a Daily Trial Balance workbook tab.

    python skills/trial-balance-reconciliation/scripts/extract_formulas.py [TAB]

Why this exists: the client asked, reasonably, whether the sheet's formulas were
ever recorded anywhere - because re-deriving them by hand is painful and has
already been done more than once. They had not been, for the Trial Balance. They
never need to be again: openpyxl reads the formulas straight out of the workbook,
so the register is *generated*, not transcribed. Re-run this against any tab and
diff it against the last one to see exactly what the station changed.

Output columns: cell · the label on that row · the formula · the value Excel
computed for it. The last column is what makes it a reconciliation source rather
than just a listing.
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# The sheet carries rupee signs and curly quotes; a Windows console is cp1252.
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

REPO = Path(__file__).resolve().parents[3]
PATH = REPO / "docs/01-BRD-Requirement-Gathering/ocr-samples/Trail_balance_12-SEP-2026.xlsx"
TAB = sys.argv[1] if len(sys.argv) > 1 else "SEP12"
MAX_ROW, MAX_COL = 140, 26


def main() -> None:
    formulas = load_workbook(PATH, data_only=False)[TAB]
    values = load_workbook(PATH, data_only=True)[TAB]

    count = 0
    for row in formulas.iter_rows(min_row=1, max_row=MAX_ROW, max_col=MAX_COL):
        # The row's own label, for reading the register without the workbook open.
        label = None
        for cell in row[:2]:
            v = values.cell(row=cell.row, column=cell.column).value
            if isinstance(v, str) and v.strip():
                label = " ".join(v.split())[:52]
                break
        for cell in row:
            formula = cell.value
            if not (isinstance(formula, str) and formula.startswith("=")):
                continue
            count += 1
            ref = f"{get_column_letter(cell.column)}{cell.row}"
            computed = values.cell(row=cell.row, column=cell.column).value
            shown = f"{computed:,.4f}" if isinstance(computed, (int, float)) else str(computed)
            print(f"{ref:>6} | {(label or ''):52} | {formula:38} | {shown}")
    print(f"\n{count} formulas in {TAB}", file=sys.stderr)


if __name__ == "__main__":
    main()
